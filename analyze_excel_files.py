"""
Excel文件格式分析工具
用于分析资产快照和操作对账单的实际格式
"""

import os
import sys

def try_read_excel(filepath):
    """尝试多种方式读取Excel文件"""
    print(f"\n分析文件: {filepath}")
    print("=" * 80)

    # 方法1: 使用pandas (xlrd引擎)
    try:
        import pandas as pd
        print("\n方法1: pandas + xlrd")
        df = pd.read_excel(filepath, engine='xlrd')
        print(f"✓ 成功读取")
        print(f"  形状: {df.shape}")
        print(f"  列名: {df.columns.tolist()}")
        print(f"  前5行:")
        print(df.head())
        print(f"  数据类型:")
        print(df.dtypes)
        return df
    except Exception as e:
        print(f"✗ 失败: {e}")

    # 方法2: 使用pandas (openpyxl引擎)
    try:
        import pandas as pd
        print("\n方法2: pandas + openpyxl")
        df = pd.read_excel(filepath, engine='openpyxl')
        print(f"✓ 成功读取")
        print(f"  形状: {df.shape}")
        print(f"  列名: {df.columns.tolist()}")
        print(f"  前5行:")
        print(df.head())
        print(f"  数据类型:")
        print(df.dtypes)
        return df
    except Exception as e:
        print(f"✗ 失败: {e}")

    # 方法3: 使用xlrd直接读取
    try:
        import xlrd
        print("\n方法3: xlrd直接读取")
        workbook = xlrd.open_workbook(filepath)
        print(f"✓ 成功打开")
        print(f"  工作表数量: {workbook.nsheets}")
        print(f"  工作表名称: {workbook.sheet_names()}")

        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            print(f"\n  工作表: {sheet_name}")
            print(f"    行数: {sheet.nrows}, 列数: {sheet.ncols}")

            # 读取前5行
            print(f"    前5行数据:")
            for row_idx in range(min(5, sheet.nrows)):
                row = sheet.row_values(row_idx)
                print(f"      行{row_idx}: {row}")

        return None
    except Exception as e:
        print(f"✗ 失败: {e}")

    # 方法4: 使用openpyxl直接读取
    try:
        import openpyxl
        print("\n方法4: openpyxl直接读取")
        workbook = openpyxl.load_workbook(filepath)
        print(f"✓ 成功打开")
        print(f"  工作表数量: {len(workbook.sheetnames)}")
        print(f"  工作表名称: {workbook.sheetnames}")

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            print(f"\n  工作表: {sheet_name}")
            print(f"    最大行数: {sheet.max_row}, 最大列数: {sheet.max_column}")

            # 读取前5行
            print(f"    前5行数据:")
            for row_idx in range(1, min(6, sheet.max_row + 1)):
                row_data = []
                for col_idx in range(1, sheet.max_column + 1):
                    cell_value = sheet.cell(row_idx, col_idx).value
                    row_data.append(str(cell_value) if cell_value is not None else '')
                print(f"      行{row_idx}: {row_data}")

        return None
    except Exception as e:
        print(f"✗ 失败: {e}")

    return None


if __name__ == '__main__':
    print("=" * 80)
    print("Excel文件格式分析工具")
    print("=" * 80)

    # 分析资产快照文件
    snapshot_dir = 'StockAssetSnapshot'
    if os.path.exists(snapshot_dir):
        print("\n" + "=" * 80)
        print("【资产快照文件】")
        print("=" * 80)

        for filename in os.listdir(snapshot_dir):
            if filename.endswith('.xls') or filename.endswith('.xlsx'):
                filepath = os.path.join(snapshot_dir, filename)
                try_read_excel(filepath)
    else:
        print(f"✗ 目录不存在: {snapshot_dir}")

    # 分析操作对账单文件
    operation_dir = 'StockOperationLog'
    if os.path.exists(operation_dir):
        print("\n" + "=" * 80)
        print("【操作对账单文件】")
        print("=" * 80)

        for filename in os.listdir(operation_dir):
            if filename.endswith('.xls') or filename.endswith('.xlsx'):
                filepath = os.path.join(operation_dir, filename)
                try_read_excel(filepath)
    else:
        print(f"✗ 目录不存在: {operation_dir}")

    print("\n" + "=" * 80)
    print("分析完成")
    print("=" * 80)
